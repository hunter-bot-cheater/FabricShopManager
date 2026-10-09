from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from app.services import inventory_service


def test_inventory_export_contains_stock_events_and_purchases(tmp_path, monkeypatch):
    target_dir = tmp_path / "exports"
    monkeypatch.setattr(inventory_service, "load_settings", lambda: {"output_inventory": str(target_dir)})
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("""CREATE TABLE app_products (
            flower TEXT, cost_per_meter REAL, roll_length_meters REAL, replenishment_days INTEGER,
            current_stock REAL, active INTEGER)"""))
        connection.execute(text("""CREATE TABLE app_stock_events (
            id INTEGER, effective_date TEXT, flower TEXT, event_type TEXT, change_qty REAL,
            before_stock REAL, after_stock REAL, reference TEXT, created_at TEXT)"""))
        connection.execute(text("""CREATE TABLE app_purchase_receipts (
            id INTEGER, received_date TEXT, flower TEXT, roll_count INTEGER, roll_length_meters REAL,
            total_meters REAL, cost_per_meter REAL, total_cost REAL, reference TEXT)"""))
        connection.execute(text("INSERT INTO app_products VALUES ('花型A', 12, 45, 0, 90, 1)"))
        connection.execute(text("INSERT INTO app_stock_events VALUES (1, '2026-10-08', '花型A', '采购入库', 45, 45, 90, '1卷', '2026-10-08 10:00:00')"))
        connection.execute(text("INSERT INTO app_purchase_receipts VALUES (1, '2026-10-08', '花型A', 1, 45, 45, 12, 540, '供应商A')"))

    path = inventory_service.export_inventory(engine)
    assert path.parent == target_dir
    assert path.exists()
    workbook = load_workbook(path, read_only=True)
    assert workbook.sheetnames == ["当前库存", "库存流水", "采购到货"]
    assert workbook["当前库存"]["A2"].value == "花型A"
    assert workbook["采购到货"]["C2"].value == 1

