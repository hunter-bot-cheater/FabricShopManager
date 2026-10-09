from datetime import date
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

from app.services import report_service


def _report_engine():
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("""CREATE TABLE app_orders (
            id INTEGER PRIMARY KEY, platform INTEGER, order_no TEXT, parent_order_no TEXT,
            product TEXT, product_spec TEXT, order_status TEXT, after_sale_status TEXT,
            product_quantity REAL, merchant_income REAL, postage REAL, flower TEXT, meters REAL,
            order_date DATE, express_no TEXT, refunded INTEGER, cancelled INTEGER)"""))
        connection.execute(text("CREATE TABLE app_products (flower TEXT, cost_per_meter REAL)"))
        connection.execute(text("INSERT INTO app_products VALUES ('花型A', 10)"))
    return engine


def _insert_order(connection, order_id, platform, order_no, express_no=None, parent_order_no=None):
    connection.execute(text("""INSERT INTO app_orders
        (id, platform, order_no, parent_order_no, product, product_spec, order_status,
         after_sale_status, product_quantity, merchant_income, postage, flower, meters,
         order_date, express_no, refunded, cancelled)
        VALUES (:id, :platform, :order_no, :parent, '布料', '1米', '已发货', '',
         1, 30, 0, '花型A', 1, :day, :express, 0, 0)"""),
        {"id": order_id, "platform": platform, "order_no": order_no, "parent": parent_order_no,
         "day": date(2026, 10, 8), "express": express_no})


def test_same_pdd_and_taobao_tracking_number_is_one_package(monkeypatch):
    monkeypatch.setattr(report_service, "load_settings", lambda: {"default_shipping_fee": 2.5})
    engine = _report_engine()
    with engine.begin() as connection:
        _insert_order(connection, 1, 0, "PDD-1", express_no="SF001")
        _insert_order(connection, 2, 1, "TB-1", express_no="SF001")

    result = report_service.build_report(engine, date(2026, 10, 8), date(2026, 10, 8))

    assert result["totals"]["运费"] == 2.5
    assert sum(row["快递费"] for row in result["details"]) == 2.5


def test_douyin_parent_key_does_not_merge_with_other_platform_tracking_number(monkeypatch):
    monkeypatch.setattr(report_service, "load_settings", lambda: {"default_shipping_fee": 2.5})
    engine = _report_engine()
    with engine.begin() as connection:
        _insert_order(connection, 1, 2, "DY-SUB-1", express_no="SF777", parent_order_no="DY999")
        _insert_order(connection, 2, 0, "PDD-2", express_no="SF777")

    result = report_service.build_report(engine, date(2026, 10, 8), date(2026, 10, 8))

    assert result["totals"]["运费"] == 5.0
    assert sum(row["快递费"] for row in result["details"]) == 5.0


def test_missing_shipping_identifier_falls_back_to_order_number(monkeypatch):
    monkeypatch.setattr(report_service, "load_settings", lambda: {"default_shipping_fee": 2.5})
    engine = _report_engine()
    with engine.begin() as connection:
        _insert_order(connection, 1, 0, "PDD-3")
        _insert_order(connection, 2, 0, "PDD-3")

    result = report_service.build_report(engine, date(2026, 10, 8), date(2026, 10, 8))

    assert result["totals"]["运费"] == 2.5


def test_business_report_excludes_unpaid_rows_even_if_status_mentions_shipped(monkeypatch):
    monkeypatch.setattr(report_service, "load_settings", lambda: {"default_shipping_fee": 2.5})
    engine = _report_engine()
    with engine.begin() as connection:
        _insert_order(connection, 1, 0, "PDD-UNPAID", express_no="SF001")
        connection.execute(text("UPDATE app_orders SET order_status='未付款，已发货' WHERE id=1"))

    result = report_service.build_report(engine, date(2026, 10, 8), date(2026, 10, 8))

    assert result["rows"] == []
    assert result["totals"]["营业额"] == 0


def test_daily_export_keeps_legacy_name_and_sheet_layout(tmp_path, monkeypatch):
    monkeypatch.setattr(report_service, "load_settings", lambda: {
        "default_shipping_fee": 2.5, "output_daily": str(tmp_path)})
    engine = _report_engine()
    with engine.begin() as connection:
        _insert_order(connection, 1, 0, "PDD-1", express_no="SF001")

    target = report_service.export_report(engine, date(2026, 10, 8), date(2026, 10, 8), "daily")

    assert target.name == "20261008日报.xlsx"
    workbook = load_workbook(target, read_only=False)
    assert workbook.sheetnames == ["花型汇总", "平台汇总", "订单明细"]
    assert workbook["花型汇总"]["A1"].value == "花型"
    assert workbook["订单明细"]["A1"].value == "花型"
    assert workbook["订单明细"]["K1"].value == "成交日期"
    assert workbook["订单明细"]["L1"].value == "商品规格"


def test_legacy_report_total_uses_order_rows_while_platform_counts_are_distinct(tmp_path, monkeypatch):
    monkeypatch.setattr(report_service, "load_settings", lambda: {
        "default_shipping_fee": 2.5, "output_daily": str(tmp_path)})
    engine = _report_engine()
    with engine.begin() as connection:
        _insert_order(connection, 1, 0, "PDD-1", express_no="SF001")
        _insert_order(connection, 2, 0, "PDD-1", express_no="SF001")

    target = report_service.export_report(engine, date(2026, 10, 8), date(2026, 10, 8), "daily")

    workbook = load_workbook(target, read_only=False)
    assert workbook["平台汇总"]["B2"].value == 1
    assert workbook["平台汇总"]["B5"].value == 2
    summary = workbook["花型汇总"]
    total_col = next(cell.column for cell in summary[2] if cell.value == "订单数" and cell.column > 20)
    assert summary.cell(summary.max_row, total_col).value == 2


def test_inout_export_uses_legacy_shape_and_reconciles_stock(tmp_path, monkeypatch):
    monkeypatch.setattr(report_service, "load_settings", lambda: {"output_inout": str(tmp_path)})
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE app_products (flower TEXT, current_stock REAL, active INTEGER)"))
        connection.execute(text("CREATE TABLE app_meta (meta_key TEXT, meta_value TEXT)"))
        connection.execute(text("""CREATE TABLE app_stock_events (
            id INTEGER PRIMARY KEY, movement_key TEXT, flower TEXT, effective_date DATE, event_type TEXT,
            change_qty REAL, before_stock REAL, after_stock REAL, reference TEXT, created_at DATETIME)"""))
        connection.execute(text("INSERT INTO app_products VALUES ('花型A', 90, 1)"))
        connection.execute(text("INSERT INTO app_meta VALUES ('legacy_baseline_v1', '2026-10-01T09:00:00')"))
        connection.execute(text("""INSERT INTO app_stock_events
            (movement_key, flower, effective_date, event_type, change_qty, before_stock, after_stock)
            VALUES ('opening', '花型A', '2026-10-01', '期初库存', 100, 0, 100),
                   ('out', '花型A', '2026-10-07', '订单出库', -20, 100, 80),
                   ('in', '花型A', '2026-10-08', '采购入库', 10, 80, 90)"""))

    target = report_service.export_inout_report(engine, date(2026, 10, 7), date(2026, 10, 8))

    assert target.name == "261007-261008出库入库明细.xlsx"
    sheet = load_workbook(target, read_only=False)["出库入库明细"]
    assert sheet["A1"].value == "261007-261008出库入库明细"
    assert sheet["A4"].value == 100
    assert sheet["D4"].value == 20
    assert sheet["E5"].value == 10
    assert sheet["G4"].value == 90


def test_inout_screen_and_export_share_legacy_history(tmp_path, monkeypatch):
    monkeypatch.setattr(report_service, "load_settings", lambda: {"output_inout": str(tmp_path)})
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        connection.execute(text("CREATE TABLE app_products (flower TEXT, current_stock REAL, active INTEGER)"))
        connection.execute(text("CREATE TABLE app_meta (meta_key TEXT, meta_value TEXT)"))
        connection.execute(text("""CREATE TABLE app_stock_events (
            id INTEGER PRIMARY KEY, movement_key TEXT, flower TEXT, effective_date DATE, event_type TEXT,
            change_qty REAL, before_stock REAL, after_stock REAL, reference TEXT, created_at DATETIME)"""))
        connection.execute(text("CREATE TABLE inventory_snapshot (snapshot_date DATE, flower TEXT, stock REAL)"))
        connection.execute(text("CREATE TABLE daily_report_cache (report_date DATE, flower TEXT, total_meters REAL)"))
        connection.execute(text("""CREATE TABLE inventory_log (
            effect_date DATE, created_at DATETIME, flower TEXT, change_type TEXT, change_qty REAL)"""))
        connection.execute(text("INSERT INTO app_products VALUES ('花型A', 125, 1)"))
        connection.execute(text("INSERT INTO app_meta VALUES ('legacy_baseline_v1', '2026-10-08T09:00:00')"))
        connection.execute(text("INSERT INTO inventory_snapshot VALUES ('2026-10-06', '花型A', 120)"))
        connection.execute(text("INSERT INTO daily_report_cache VALUES ('2026-10-07', '花型A', 5)"))
        connection.execute(text("INSERT INTO inventory_log VALUES ('2026-10-07', '2026-10-07 12:00:00', '花型A', '入库', 10)"))

    report = report_service.build_inout_report(engine, date(2026, 10, 7), date(2026, 10, 7))
    target = report_service.export_inout_report(engine, date(2026, 10, 7), date(2026, 10, 7))
    sheet = load_workbook(target, read_only=True)["出库入库明细"]

    assert report["rows"] == [{"花型": "花型A", "期初库存": 120, "入库米数": 10,
                              "出库米数": 5, "期末库存": 125}]
    assert {(item["event_type"], item["change_qty"]) for item in report["details"]} == {
        ("订单出库", -5), ("采购入库", 10)}
    assert sheet["A4"].value == report["rows"][0]["期初库存"]
    assert sheet["D4"].value == report["rows"][0]["出库米数"]
    assert sheet["D5"].value == report["rows"][0]["入库米数"]
    assert sheet["F4"].value == report["rows"][0]["期末库存"]
