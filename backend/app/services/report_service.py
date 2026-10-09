from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import inspect, text
from sqlalchemy.engine import Engine

from app.core.settings import load_settings
from app.domain.order_rules import PLATFORM_NAMES, is_shipped


def build_report(engine: Engine, start_date: date, end_date: date, group_by: str = "day") -> dict:
    if end_date < start_date:
        raise ValueError("结束日期不能早于开始日期")
    with engine.connect() as connection:
        rows = connection.execute(text("""SELECT id, platform, order_no, parent_order_no, product, product_spec,
            order_status, after_sale_status, product_quantity, merchant_income, postage, flower, meters,
            order_date, express_no, refunded, cancelled FROM app_orders
            WHERE order_date BETWEEN :start AND :end AND flower IS NOT NULL
            ORDER BY id"""), {"start": start_date, "end": end_date}).mappings().all()
        costs = {row["flower"]: float(row["cost_per_meter"] or 0) for row in connection.execute(
            text("SELECT flower, cost_per_meter FROM app_products")).mappings()}

    eligible = []
    settings = load_settings()
    for source in rows:
        row = dict(source)
        if row["refunded"] or row["cancelled"] or not is_shipped(row):
            continue
        if row["flower"] not in costs:
            continue
        row["meters"] = float(row["meters"] or 0)
        row["revenue"] = float(row["merchant_income"] or 0)
        row["cost"] = round(row["meters"] * costs[row["flower"]], 2)
        row["shipping_fee"] = 0.0
        row["platform_name"] = PLATFORM_NAMES.get(int(row["platform"] or 0), "拼多多")
        # Match the legacy order_key: Douyin uses its parent order number;
        # PDD/Taobao use the tracking number; missing values fall back to order_no.
        primary = row["parent_order_no"] if int(row["platform"] or 0) == 2 else row["express_no"]
        primary = "" if primary is None or pd.isna(primary) else str(primary).strip()
        fallback = row["order_no"]
        fallback = "" if fallback is None or pd.isna(fallback) else str(fallback).strip()
        row["package_key"] = primary or fallback
        eligible.append(row)

    package_rows: dict[str, list[dict]] = defaultdict(list)
    for row in eligible:
        package_rows[str(row["package_key"])].append(row)
    for package in package_rows.values():
        first = package[0]
        first["shipping_fee"] = float(settings["default_shipping_fee"])

    package_numbers = {}
    for row in eligible:
        if row["package_key"] not in package_numbers:
            package_numbers[row["package_key"]] = len(package_numbers) + 1

    summary: dict[tuple[str, str, str], dict] = {}
    detail = []
    for row in eligible:
        order_date = row["order_date"]
        if isinstance(order_date, str):
            order_date = date.fromisoformat(order_date[:10])
        bucket = order_date.strftime("%Y-%m") if group_by == "month" else order_date.isoformat()
        key = (bucket, row["flower"], row["platform_name"])
        values = summary.setdefault(key, {"日期": bucket, "花型": row["flower"], "平台": row["platform_name"],
                                          "订单行数": 0, "米数": 0.0, "营业额": 0.0, "成本": 0.0,
                                          "运费": 0.0, "利润": 0.0})
        values["订单行数"] += 1
        values["米数"] += row["meters"]
        values["营业额"] += row["revenue"]
        values["成本"] += row["cost"]
        values["运费"] += row["shipping_fee"]
        values["利润"] += row["revenue"] - row["cost"] - row["shipping_fee"]
        detail.append({"日期": bucket, "平台": row["platform_name"], "发货包号": package_numbers[row["package_key"]],
                       "订单号": row["order_no"], "商品": row["product"], "规格": row["product_spec"],
                       "商品数量": float(row["product_quantity"] or 0), "售后状态": row["after_sale_status"],
                       "是否退款": False, "花型": row["flower"], "米数": row["meters"],
                       "营业额": row["revenue"], "成本": row["cost"], "快递费": row["shipping_fee"],
                       "盈利": row["revenue"] - row["cost"] - row["shipping_fee"]})

    data = sorted(summary.values(), key=lambda item: (item["日期"], item["花型"], item["平台"]))
    for row in data:
        for key in ("米数", "营业额", "成本", "运费", "利润"):
            row[key] = round(row[key], 2)
    totals = {"订单行数": sum(row["订单行数"] for row in data), "米数": round(sum(row["米数"] for row in data), 2),
              "营业额": round(sum(row["营业额"] for row in data), 2), "成本": round(sum(row["成本"] for row in data), 2),
              "运费": round(sum(row["运费"] for row in data), 2), "利润": round(sum(row["利润"] for row in data), 2)}
    return {"rows": data, "details": detail, "totals": totals, "start_date": start_date.isoformat(),
            "end_date": end_date.isoformat(), "group_by": group_by}


def build_refund_report(engine: Engine, start_date: date, end_date: date) -> dict:
    if end_date < start_date:
        raise ValueError("结束日期不能早于开始日期")
    with engine.connect() as connection:
        source_rows = connection.execute(text("""SELECT platform, order_no, product, product_spec,
            flower, order_status, after_sale_status, merchant_income, order_date, delivery_time,
            receive_time, refunded, cancelled FROM app_orders
            WHERE order_date BETWEEN :start AND :end ORDER BY id"""),
            {"start": start_date, "end": end_date}).mappings().all()

    platforms = list(PLATFORM_NAMES.values())
    refunds: dict[tuple[str, str], float] = defaultdict(float)
    turnover: dict[tuple[str, str], float] = defaultdict(float)
    details = []
    for source in source_rows:
        row = dict(source)
        platform = PLATFORM_NAMES.get(int(row["platform"] or 0), "拼多多")
        flower = row["flower"] or "未匹配花型"
        shipped = is_shipped(row) or bool(row["delivery_time"] or row["receive_time"])
        if row["cancelled"] or not shipped:
            continue
        amount = float(row["merchant_income"] or 0)
        if row["refunded"]:
            refunds[(platform, flower)] += amount
            details.append({"平台": platform, "花型": flower, "订单号": row["order_no"],
                            "商品": row["product"], "规格": row["product_spec"],
                            "订单金额": amount, "业务日期": str(row["order_date"] or "")})
        else:
            turnover[(platform, flower)] += amount

    platform_totals = {}
    all_refunds = all_turnover = 0.0
    for platform in platforms:
        refund_amount = sum(v for (p, _), v in refunds.items() if p == platform)
        base_amount = sum(v for (p, _), v in turnover.items() if p == platform)
        all_refunds += refund_amount
        all_turnover += base_amount
        platform_totals[platform] = {"退款金额": round(refund_amount, 2), "退款基数": round(base_amount, 2),
                                     "退款率": round(refund_amount / base_amount, 6) if base_amount else 0.0}

    flowers = sorted({flower for _, flower in set(refunds) | set(turnover)})
    report_rows = []
    for flower in flowers:
        row = {"花型": flower}
        total_refunds = total_turnover = 0.0
        for platform in platforms:
            amount, base = refunds[(platform, flower)], turnover[(platform, flower)]
            row[f"{platform}_退款金额"] = round(amount, 2)
            row[f"{platform}_退款基数"] = round(base, 2)
            row[f"{platform}_退款率"] = round(amount / base, 6) if base else 0.0
            total_refunds += amount
            total_turnover += base
        row["汇总_退款金额"] = round(total_refunds, 2)
        row["汇总_退款基数"] = round(total_turnover, 2)
        row["汇总_退款率"] = round(total_refunds / total_turnover, 6) if total_turnover else 0.0
        report_rows.append(row)

    totals = {"退款金额": round(all_refunds, 2), "退款基数": round(all_turnover, 2),
              "退款率": round(all_refunds / all_turnover, 6) if all_turnover else 0.0}
    return {"rows": report_rows, "details": details, "platforms": platforms,
            "platform_totals": platform_totals, "totals": totals,
            "start_date": start_date.isoformat(), "end_date": end_date.isoformat()}


def export_refund_report(engine: Engine, start_date: date, end_date: date) -> Path:
    report = build_refund_report(engine, start_date, end_date)
    folder = Path(load_settings()["output_refund"])
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"退款明细_{start_date:%Y%m%d}-{end_date:%Y%m%d}.xlsx"
    columns = ["花型"]
    for platform in report["platforms"]:
        columns.extend((f"{platform}_退款金额", f"{platform}_退款基数", f"{platform}_退款率"))
    columns.extend(("汇总_退款金额", "汇总_退款基数", "汇总_退款率"))
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        pd.DataFrame(report["rows"], columns=columns).to_excel(writer, sheet_name="花型退款汇总", index=False)
        pd.DataFrame([{"平台": name, **values} for name, values in report["platform_totals"].items()] +
                     [{"平台": "汇总", **report["totals"]}]).to_excel(writer, sheet_name="平台退款汇总", index=False)
        pd.DataFrame(report["details"], columns=("平台", "花型", "订单号", "商品", "规格", "订单金额", "业务日期"))\
            .to_excel(writer, sheet_name="已发货退款明细", index=False)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(max(max((len(str(cell.value or "")) for cell in column), default=0) + 2, 10), 36)
                sheet.column_dimensions[column[0].column_letter].width = width
    return target


def _inout_data(engine: Engine, start_date: date, end_date: date) -> dict:
    days = [start_date + timedelta(days=offset) for offset in range((end_date - start_date).days + 1)]
    daily_out, daily_in = defaultdict(float), defaultdict(float)
    details = []
    with engine.connect() as connection:
        products = connection.execute(text(
            "SELECT flower, current_stock FROM app_products WHERE active=1 ORDER BY flower"
        )).mappings().all()
        table_names = set(inspect(connection).get_table_names())
        marker = connection.execute(text(
            "SELECT meta_value FROM app_meta WHERE meta_key='legacy_baseline_v1'"
        )).scalar() if "app_meta" in table_names else None
        baseline = date.fromisoformat(str(marker)[:10]) if marker else None
        use_legacy = baseline is not None and start_date < baseline
        openings = {}
        if use_legacy:
            required = {"inventory_snapshot", "daily_report_cache", "inventory_log"}
            if not required.issubset(table_names):
                raise ValueError("所选日期早于新系统库存基线，但旧版库存快照/流水表不可用，无法准确生成报表")
            legacy_end = min(end_date, baseline - timedelta(days=1))
            previous = start_date - timedelta(days=1)
            openings.update({row[0]: float(row[1] or 0) for row in connection.execute(text(
                "SELECT flower, stock FROM inventory_snapshot WHERE snapshot_date=:day"),
                {"day": previous}).all()})
            if legacy_end >= start_date:
                legacy_outflows = []
                has_imported_legacy_orders = False
                if "app_orders" in table_names:
                    has_imported_legacy_orders = connection.execute(text(
                        "SELECT COUNT(*) FROM app_orders WHERE legacy=1"
                    )).scalar_one() > 0
                    legacy_outflows = connection.execute(text("""SELECT inventory_date AS report_date,
                        flower, SUM(baseline_stock_effect_meters) AS total_meters
                        FROM app_orders WHERE legacy=1 AND flower IS NOT NULL
                        AND inventory_date BETWEEN :start AND :end
                        GROUP BY inventory_date, flower"""),
                        {"start": start_date, "end": legacy_end}).all()
                if not has_imported_legacy_orders:
                    legacy_outflows = connection.execute(text("""SELECT report_date, flower, total_meters
                        FROM daily_report_cache WHERE report_date BETWEEN :start AND :end"""),
                        {"start": start_date, "end": legacy_end}).all()
                for row in legacy_outflows:
                    day, flower, quantity = _report_date(row[0]), row[1], float(row[2] or 0)
                    daily_out[(day, flower)] += quantity
                    if quantity:
                        details.append({"id": None, "flower": flower, "effective_date": day,
                            "event_type": "订单出库", "change_qty": -quantity, "before_stock": None,
                            "after_stock": None, "reference": "旧系统订单基线出库", "created_at": None})
                for row in connection.execute(text("""SELECT COALESCE(effect_date, DATE(created_at)) AS day,
                    flower, SUM(change_qty) FROM inventory_log
                    WHERE change_type='入库' AND COALESCE(effect_date, DATE(created_at)) BETWEEN :start AND :end
                    GROUP BY day, flower"""), {"start": start_date, "end": legacy_end}).all():
                    day, flower, quantity = _report_date(row[0]), row[1], float(row[2] or 0)
                    daily_in[(day, flower)] += quantity
                    if quantity:
                        details.append({"id": None, "flower": flower, "effective_date": day,
                            "event_type": "采购入库", "change_qty": quantity, "before_stock": None,
                            "after_stock": None, "reference": "旧系统库存流水历史入库", "created_at": None})
            modern_start = max(start_date, baseline)
        else:
            modern_start = start_date

        events = connection.execute(text("""SELECT id, flower, effective_date, event_type, change_qty,
            before_stock, after_stock, reference, created_at FROM app_stock_events
            WHERE effective_date BETWEEN :start AND :end AND event_type<>'期初库存'
            ORDER BY effective_date, id"""), {"start": modern_start, "end": end_date}).mappings().all()
        changes = connection.execute(text("""SELECT flower, SUM(change_qty) FROM app_stock_events
            WHERE effective_date >= :start AND event_type <> '期初库存' GROUP BY flower"""),
            {"start": modern_start}).all()
        future_delta = {row[0]: float(row[1] or 0) for row in changes}
        if not use_legacy:
            openings = {row["flower"]: float(row["current_stock"] or 0) - future_delta.get(row["flower"], 0)
                        for row in products}
        for event in events:
            item = dict(event)
            details.append(item)
            quantity = float(event["change_qty"] or 0)
            key = (_report_date(event["effective_date"]), event["flower"])
            if quantity < 0:
                daily_out[key] += -quantity
            elif quantity > 0:
                daily_in[key] += quantity
    return {"days": days, "products": products, "openings": openings,
            "daily_out": daily_out, "daily_in": daily_in, "details": details}


def build_inout_report(engine: Engine, start_date: date, end_date: date) -> dict:
    if end_date < start_date:
        raise ValueError("结束日期不能早于开始日期")
    data = _inout_data(engine, start_date, end_date)
    rows = []
    for product in data["products"]:
        flower = product["flower"]
        opening = data["openings"].get(flower, 0.0)
        incoming = sum(data["daily_in"][(day, flower)] for day in data["days"])
        outgoing = sum(data["daily_out"][(day, flower)] for day in data["days"])
        rows.append({"花型": flower, "期初库存": round(opening, 2),
                     "入库米数": round(incoming, 2), "出库米数": round(outgoing, 2),
                     "期末库存": round(opening + incoming - outgoing, 2)})
    return {"rows": rows, "details": data["details"],
            "start_date": start_date.isoformat(), "end_date": end_date.isoformat()}


def export_report(engine: Engine, start_date: date, end_date: date, kind: str) -> Path:
    if kind not in {"daily", "range"}:
        raise ValueError("导出类型不支持")
    if kind == "daily" and start_date != end_date:
        raise ValueError("日报只能选择同一天；多日统计请使用区间报告")
    report = build_report(engine, start_date, end_date, "day")
    settings = load_settings()
    folder_key = {"daily": "output_daily", "range": "output_range"}[kind]
    folder = Path(settings[folder_key])
    folder.mkdir(parents=True, exist_ok=True)
    if kind == "daily":
        target = folder / f"{start_date:%Y%m%d}日报.xlsx"
    elif kind == "range":
        target = folder / f"区间报告_{start_date:%Y%m%d}-{end_date:%Y%m%d}.xlsx"
    return _export_original_daily_layout(target, report)


REPORT_METRICS = ("订单数", "成本", "米数", "营业额", "快递费", "盈利")
DETAIL_COLUMNS = ("花型", "平台", "发货包号", "成本", "米数", "营业额", "快递费", "盈利",
                  "售后状态", "订单号", "成交日期", "商品规格", "商品数量", "是否退款")


def _export_original_daily_layout(target: Path, report: dict) -> Path:
    platform_names = list(PLATFORM_NAMES.values())
    source = report["details"]
    by_flower: dict[str, dict] = defaultdict(lambda: {p: defaultdict(float) for p in platform_names} | {"_orders": {p: set() for p in platform_names}})
    platform_totals = {p: defaultdict(float) for p in platform_names}
    platform_orders = {p: set() for p in platform_names}
    for item in source:
        flower, platform = item["花型"], item["平台"]
        if platform not in platform_names:
            continue
        values = by_flower[flower][platform]
        values["成本"] += item["成本"]
        values["米数"] += item["米数"]
        values["营业额"] += item["营业额"]
        values["快递费"] += item["快递费"]
        values["盈利"] += item["盈利"]
        by_flower[flower]["_orders"][platform].add(item["订单号"])
        for key in ("成本", "米数", "营业额", "快递费", "盈利"):
            platform_totals[platform][key] += item[key]
        platform_orders[platform].add(item["订单号"])

    wide_columns = ["花型"]
    for index, name in enumerate(platform_names):
        wide_columns.extend(f"{name}_{metric}" for metric in REPORT_METRICS)
        wide_columns.append(f"_spacer_{index}")
    wide_columns.extend(f"汇总_{metric}" for metric in REPORT_METRICS)
    summary_rows = []
    flowers = sorted(by_flower, key=lambda name: sum(by_flower[name][p]["营业额"] for p in platform_names), reverse=True)
    for flower in flowers:
        row = {"花型": flower}
        total = defaultdict(float)
        for index, name in enumerate(platform_names):
            values = by_flower[flower][name]
            row[f"{name}_订单数"] = len(by_flower[flower]["_orders"][name])
            for metric in REPORT_METRICS[1:]:
                row[f"{name}_{metric}"] = round(values[metric], 2)
                total[metric] += values[metric]
            row[f"_spacer_{index}"] = ""
            total["订单数"] += len(by_flower[flower]["_orders"][name])
        row.update({f"汇总_{key}": round(value, 2) if key != "订单数" else int(value) for key, value in total.items()})
        summary_rows.append(row)
    totals_row = {"花型": "【总计】"}
    for index, name in enumerate(platform_names):
        totals_row[f"{name}_订单数"] = len(platform_orders[name])
        for metric in REPORT_METRICS[1:]:
            totals_row[f"{name}_{metric}"] = round(platform_totals[name][metric], 2)
        totals_row[f"_spacer_{index}"] = ""
    totals_row["汇总_订单数"] = len(source)
    for metric in REPORT_METRICS[1:]:
        totals_row[f"汇总_{metric}"] = round(sum(platform_totals[p][metric] for p in platform_names), 2)
    summary_rows.append(totals_row)

    platform_rows = []
    for name in platform_names:
        values = platform_totals[name]
        platform_rows.append({"平台": name, "订单数": len(platform_orders[name]),
                              **{key: round(values[key], 2) for key in REPORT_METRICS[1:]}})
    platform_rows.append({"平台": "合计", "订单数": len(source),
                          **{key: round(sum(platform_totals[p][key] for p in platform_names), 2)
                             for key in REPORT_METRICS[1:]}})

    details = []
    for item in sorted(source, key=lambda row: (row["花型"], row["订单号"])):
        details.append({"花型": item["花型"], "平台": item["平台"], "发货包号": item["发货包号"],
                        "成本": item["成本"], "米数": item["米数"], "营业额": item["营业额"],
                        "快递费": item["快递费"], "盈利": item["盈利"], "售后状态": item["售后状态"],
                        "订单号": item["订单号"], "成交日期": item["日期"], "商品规格": item["规格"],
                        "商品数量": item["商品数量"], "是否退款": item["是否退款"]})

    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        pd.DataFrame(summary_rows, columns=wide_columns).to_excel(
            writer, sheet_name="花型汇总", index=False, header=False, startrow=2)
        ws = writer.sheets["花型汇总"]
        header_fill = PatternFill("solid", fgColor="D9E1F2")
        border_side = __import__("openpyxl").styles.Side(style="thin", color="000000")
        border = __import__("openpyxl").styles.Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        center = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        cell = ws.cell(1, 1, "花型")
        cell.font = Font(bold=True); cell.alignment = center; cell.fill = header_fill; cell.border = border
        col = 1
        for index, name in enumerate(platform_names):
            start = col + 1
            end = start + len(REPORT_METRICS) - 1
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            top = ws.cell(1, start, name)
            top.font = Font(bold=True); top.alignment = center; top.fill = header_fill
            for j, metric in enumerate(REPORT_METRICS, start):
                ws.cell(2, j, metric).font = Font(bold=True)
                ws.cell(2, j).alignment = center; ws.cell(2, j).fill = header_fill
            for j in range(start, end + 2):
                ws.cell(1, j).border = border
                if j > end:
                    ws.merge_cells(start_row=1, start_column=j, end_row=1, end_column=j)
                    ws.cell(1, j).border = border
                    ws.cell(2, j).border = border
            col = end + 1
        start = col + 1
        end = start + len(REPORT_METRICS) - 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        top = ws.cell(1, start, "汇总")
        top.font = Font(bold=True); top.alignment = center; top.fill = header_fill
        for j, metric in enumerate(REPORT_METRICS, start):
            ws.cell(2, j, metric).font = Font(bold=True)
            ws.cell(2, j).alignment = center; ws.cell(2, j).fill = header_fill
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                cell.border = border
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        pd.DataFrame(platform_rows).to_excel(writer, sheet_name="平台汇总", index=False)
        pd.DataFrame(details, columns=DETAIL_COLUMNS).to_excel(writer, sheet_name="订单明细", index=False)
        for sheet in writer.book.worksheets:
            for index, column in enumerate(sheet.columns, start=1):
                width = min(max(max((len(str(cell.value or "")) for cell in column), default=0) + 2, 8), 30)
                sheet.column_dimensions[get_column_letter(index)].width = width
    return target


def export_inout_report(engine: Engine, start_date: date, end_date: date) -> Path:
    if end_date < start_date:
        raise ValueError("结束日期不能早于开始日期")
    settings = load_settings()
    folder = Path(settings["output_inout"])
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{start_date:%y%m%d}-{end_date:%y%m%d}出库入库明细.xlsx"
    data = _inout_data(engine, start_date, end_date)
    days, products = data["days"], data["products"]
    daily_out, daily_in, openings = data["daily_out"], data["daily_in"], data["openings"]
    if not products:
        raise ValueError("请先添加花型")
    movements_total = sum(daily_out.values()) + sum(daily_in.values())
    if movements_total == 0:
        raise ValueError("所选日期范围内无出入库流水")

    blocks = []
    for product in products:
        flower = product["flower"]
        outs = [daily_out[(day, flower)] for day in days]
        ins = [daily_in[(day, flower)] for day in days]
        out_total, in_total = sum(outs), sum(ins)
        opening = openings.get(flower, 0.0)
        blocks.append((flower, opening, outs, ins, out_total, in_total, opening + in_total - out_total))

    ncols = len(days) + 5
    last_col = get_column_letter(ncols)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "出库入库明细"
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    sheet.merge_cells(f"A1:{last_col}1")
    sheet["A1"] = f"{start_date:%y%m%d}-{end_date:%y%m%d}出库入库明细"
    sheet["A1"].font = Font(bold=True, size=14); sheet["A1"].alignment = center
    sheet.row_dimensions[1].height = 30
    sheet["B2"] = "日期"; sheet["B2"].alignment = center
    for index, day in enumerate(days):
        sheet.cell(2, 4 + index, day.day).alignment = center
    prior = start_date - timedelta(days=1)
    for col, value in ((1, f"{prior.month}月{prior.day}日库存"), (2, "花型名称"), (ncols - 1, "总出入库"), (ncols, "库存余额")):
        cell = sheet.cell(3, col, value); cell.font = Font(bold=True); cell.alignment = center; cell.fill = header_fill
    row = 4
    for flower, opening, outs, ins, out_total, in_total, balance in blocks:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=1)
        sheet.cell(row, 1, opening).number_format = "0.0"; sheet.cell(row, 1).alignment = right
        sheet.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=2)
        sheet.cell(row, 2, flower).alignment = center
        sheet.cell(row, 3, "出库米数").alignment = center; sheet.cell(row + 1, 3, "入库米数").alignment = center
        sheet.cell(row, 3).font = Font(color="FF0000")
        for index, (out_qty, in_qty) in enumerate(zip(outs, ins)):
            a = sheet.cell(row, 4 + index, out_qty); a.number_format = "0.0"; a.alignment = right; a.font = Font(color="FF0000")
            b = sheet.cell(row + 1, 4 + index, in_qty); b.number_format = "0.0"; b.alignment = right
        for target_row, total in ((row, out_total), (row + 1, in_total)):
            cell = sheet.cell(target_row, ncols - 1, total); cell.number_format = "0.0"; cell.font = Font(bold=True); cell.alignment = right
        sheet.merge_cells(start_row=row, start_column=ncols, end_row=row + 1, end_column=ncols)
        cell = sheet.cell(row, ncols, balance); cell.number_format = "0.0"; cell.font = Font(bold=True); cell.alignment = right
        row += 2
    sheet.column_dimensions["A"].width = 11; sheet.column_dimensions["B"].width = 12; sheet.column_dimensions["C"].width = 8
    for index in range(len(days)):
        sheet.column_dimensions[get_column_letter(4 + index)].width = 6
    sheet.column_dimensions[get_column_letter(ncols - 1)].width = 8; sheet.column_dimensions[get_column_letter(ncols)].width = 10
    workbook.save(target)
    return target


def _report_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else date.fromisoformat(str(value)[:10])

